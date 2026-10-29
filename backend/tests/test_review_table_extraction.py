"""Tabular review extraction: schema-driven excerpts, quote verification, export."""

import io
import json
from unittest.mock import AsyncMock, MagicMock

import pytest
from app.models.review_tables import CELL_ERROR, CELL_EXTRACTED, CELL_NOT_FOUND, CELL_UNVERIFIED
from app.services.review_tables.export import (
    CITATIONS_SHEET,
    RESULTS_SHEET,
    render_csv,
    render_xlsx,
)
from app.services.review_tables.extraction import (
    ExtractionError,
    build_prompt,
    estimate_input_tokens,
    extract_fields,
    select_excerpts,
)
from openpyxl import load_workbook

CONTRACT = (
    'This Agreement is effective as of January 5, 2026 (the "Effective Date").\n'
    "GOVERNING LAW. This Agreement shall be governed by the laws of the State of Delaware.\n"
    "LIMITATION OF LIABILITY. Liability is capped at the fees paid in the twelve months "
    "preceding the claim.\n"
)

FIELDS = [
    {"key": "governing_law", "name": "Governing law", "description": "", "type": "text"},
    {"key": "effective_date", "name": "Effective date", "description": "", "type": "date"},
    {"key": "liability_cap", "name": "Liability cap", "description": "cap amount", "type": "money"},
]


def _client_with(payload: dict) -> MagicMock:
    client = MagicMock()
    response = MagicMock()
    response.choices = [MagicMock()]
    response.choices[0].message.content = json.dumps(payload)
    client.chat.completions.create = AsyncMock(return_value=response)
    return client


class TestSelectExcerpts:
    def test_short_document_is_sent_whole(self):
        assert select_excerpts(CONTRACT, FIELDS, []) == [("document", CONTRACT)]

    def test_long_document_gets_preamble_tags_and_relevant_passages(self):
        filler = "\n\n".join(
            f"Section {i}. Boilerplate about nothing in particular." for i in range(900)
        )
        target = "INSURANCE. The Supplier shall maintain insurance coverage of USD 5,000,000."
        text = "PREAMBLE. " + filler + "\n\n" + target + "\n\n" + filler
        assert len(text) > 30_000
        start = text.index("INSURANCE.")
        tags = [{"canonical_slug": "indemnification", "span_start": 20, "span_end": 120}]
        fields = [{"key": "insurance", "name": "Insurance", "description": "coverage amount"}]

        excerpts = select_excerpts(text, fields, tags)
        labels = [label for label, _ in excerpts]
        assert labels[0] == "preamble"
        assert "indemnification" in labels
        assert any(target in snippet for _label, snippet in excerpts)
        assert sum(len(s) for _l, s in excerpts) <= 30_000 + 4_000
        assert excerpts[0][1] == text[:3_000]
        assert start > 3_000  # the target was not simply inside the preamble


class TestExtractFields:
    @pytest.mark.asyncio
    async def test_verified_quote_gets_real_offsets(self):
        quote = "This Agreement shall be governed by the laws of the State of Delaware."
        payload = {
            "fields": {
                "governing_law": {"value": "Delaware", "quote": quote, "not_found": False},
                "effective_date": {"value": "", "not_found": True},
                "liability_cap": {
                    "value": "Fees paid in prior 12 months",
                    "quote": "Liability is capped at something the model made up.",
                    "not_found": False,
                },
            }
        }
        result = await extract_fields(_client_with(payload), CONTRACT, FIELDS, model="m")

        gl = result["governing_law"]
        assert gl.status == CELL_EXTRACTED and gl.verified is True
        assert CONTRACT[gl.span_start : gl.span_end] == quote

        assert result["effective_date"].status == CELL_NOT_FOUND
        assert result["effective_date"].value is None

        cap = result["liability_cap"]
        assert cap.status == CELL_UNVERIFIED and cap.verified is False
        assert cap.value == "Fees paid in prior 12 months"
        assert cap.span_start is None

    @pytest.mark.asyncio
    async def test_missing_field_in_answer_is_not_found(self):
        payload = {
            "fields": {"governing_law": {"value": "Delaware", "quote": "", "not_found": False}}
        }
        result = await extract_fields(_client_with(payload), CONTRACT, FIELDS, model="m")
        assert set(result) == {"governing_law", "effective_date", "liability_cap"}
        assert result["effective_date"].status == CELL_NOT_FOUND
        # A value with no quote at all is kept but cannot be verified.
        assert result["governing_law"].status == CELL_UNVERIFIED

    @pytest.mark.asyncio
    async def test_provider_failure_raises_extraction_error(self):
        client = MagicMock()
        client.chat.completions.create = AsyncMock(side_effect=RuntimeError("boom"))
        with pytest.raises(ExtractionError, match="RuntimeError: boom"):
            await extract_fields(client, CONTRACT, FIELDS, model="m")

    @pytest.mark.asyncio
    async def test_unusable_answer_raises_extraction_error(self):
        client = _client_with({"nope": 1})
        with pytest.raises(ExtractionError, match="no 'fields'"):
            await extract_fields(client, CONTRACT, FIELDS, model="m")

    def test_prompt_names_every_field_with_its_type_guidance(self):
        prompt = build_prompt(FIELDS, [("document", CONTRACT)])
        assert '"governing_law"' in prompt and '"liability_cap"' in prompt
        assert "ISO date" in prompt and "currency" in prompt
        assert "character for character" in prompt


def test_estimate_is_bounded_by_the_excerpt_budget():
    small = estimate_input_tokens(4_000, 3)
    huge = estimate_input_tokens(4_000_000, 3)
    assert small < huge
    assert huge == 30_000 // 4 + 350 + 40 * 3


class TestExport:
    DOCS = [{"id": "d1", "filename": "a.pdf"}, {"id": "d2", "filename": "b.pdf"}]
    CELLS = [
        {
            "document_id": "d1",
            "field_key": "governing_law",
            "value": "Delaware",
            "quote": "governed by the laws of the State of Delaware",
            "span_start": 90,
            "span_end": 135,
            "verified": True,
            "status": CELL_EXTRACTED,
            "error": None,
        },
        {
            "document_id": "d1",
            "field_key": "effective_date",
            "value": None,
            "quote": None,
            "span_start": None,
            "span_end": None,
            "verified": False,
            "status": CELL_NOT_FOUND,
            "error": None,
        },
        {
            "document_id": "d2",
            "field_key": "governing_law",
            "value": "New York",
            "quote": "made up",
            "span_start": None,
            "span_end": None,
            "verified": False,
            "status": CELL_UNVERIFIED,
            "error": None,
        },
        {
            "document_id": "d2",
            "field_key": "effective_date",
            "value": None,
            "quote": None,
            "span_start": None,
            "span_end": None,
            "verified": False,
            "status": CELL_ERROR,
            "error": "index unavailable",
        },
    ]
    FIELDS = FIELDS[:2]

    def test_xlsx_has_results_linked_to_citations(self):
        data = render_xlsx("NDA review", self.FIELDS, self.DOCS, self.CELLS)
        wb = load_workbook(io.BytesIO(data))
        assert wb.sheetnames == [RESULTS_SHEET, CITATIONS_SHEET]

        results, citations = wb[RESULTS_SHEET], wb[CITATIONS_SHEET]
        assert [c.value for c in results[1]] == ["Document", "Governing law", "Effective date"]
        assert results["A2"].value == "a.pdf" and results["B2"].value == "Delaware"
        assert results["C2"].value == "—"  # not found: no link, no citation row
        assert results["C2"].hyperlink is None
        assert results["B2"].hyperlink.location == f"'{CITATIONS_SHEET}'!A2"
        assert results["B3"].hyperlink.location == f"'{CITATIONS_SHEET}'!A3"
        assert results["C3"].value == "error"

        rows = list(citations.iter_rows(min_row=2, values_only=True))
        assert rows[0][:4] == (
            "a.pdf",
            "Governing law",
            "Delaware",
            "governed by the laws of the State of Delaware",
        )
        assert rows[0][4:8] == (90, 135, "yes", "verified")
        assert rows[1][6:8] == ("no", "unverified")
        assert rows[2][7:9] == ("error", "index unavailable")

    def test_csv_has_one_row_per_cell(self):
        text = render_csv(self.FIELDS, self.DOCS, self.CELLS)
        lines = text.strip().splitlines()
        assert lines[0].startswith("Document,Field,Value,Quote,Span start")
        assert len(lines) == 1 + len(self.DOCS) * len(self.FIELDS)
        assert (
            "a.pdf,Governing law,Delaware,governed by the laws of the State of Delaware,90,135,yes,verified,"
            in text
        )
