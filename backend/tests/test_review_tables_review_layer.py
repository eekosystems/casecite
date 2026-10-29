"""HTTP tests for the tabular-review review layer: overrides, approvals, flags
and notes on cells and how a re-run treats them; bulk review; saved and
built-in schemas; adding and removing documents after creation; the run
history; the export options; and the source-text endpoint.
"""

import os

os.environ.setdefault("SECRET_KEY", "test-secret-key-for-testing-only-32chars!")
os.environ.setdefault("ENCRYPTION_SALT", "test-salt-16chars!")
os.environ.setdefault("DEBUG", "true")

import io
import json
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from app.models.review_tables import ReviewCell, ReviewRun
from app.services.job_queue import RetryPolicy
from app.services.review_tables.service import ReviewTableService, review_table_service
from openpyxl import load_workbook

from tests.helpers import db_count, headers, make_shared_matter, make_user, register_document
from tests.test_review_tables_api import (
    CONTRACT,
    FIELDS,
    QUOTE,
    _answer,
    _create,
    _fake_model_response,
    _poll,
    _run_to_completion,
)


@pytest.fixture(autouse=True)
def _no_rate_limit():
    with patch(
        "app.middleware.security.RateLimiter.is_allowed",
        return_value=(True, {"limit": 100, "remaining": 99}),
    ):
        yield


def _cells(client, user, tid: str) -> dict:
    detail = client.get(f"/api/v1/review-tables/{tid}", headers=headers(user)).json()
    return {(c["document_id"], c["field_key"]): c for c in detail["cells"]}


def _review(client, user, tid, doc, key, body):
    return client.patch(
        f"/api/v1/review-tables/{tid}/cells/{doc}/{key}", headers=headers(user), json=body
    )


# ------------------------------------------------------------------ one cell


def test_override_approve_and_note_one_cell(client, audit_events):
    owner = make_user("owner")
    doc = register_document(owner)
    tid = _create(client, owner, document_ids=[doc])["id"]
    _run_to_completion(client, owner, tid)

    resp = _review(
        client, owner, tid, doc, "governing_law",
        {"override_value": " New York ", "review_status": "approved", "note": "Checked §12"},
    )  # fmt: skip
    assert resp.status_code == 200, resp.text
    cell = resp.json()
    assert cell["value"] == "Delaware"  # the model's value is kept
    assert cell["override_value"] == "New York"
    assert cell["effective_value"] == "New York"
    assert cell["review_status"] == "approved"
    assert cell["note"] == "Checked §12"
    assert cell["reviewed_by"] == owner.id and cell["reviewed_at"]
    kinds = [h["kind"] for h in cell["history"]]
    assert kinds == ["model", "override", "approve"]
    assert all(h["by"] for h in cell["history"])

    # Clearing the override keeps the approval; the note may be emptied.
    resp = _review(client, owner, tid, doc, "governing_law", {"override_value": None, "note": ""})
    assert resp.status_code == 200
    cell = resp.json()
    assert cell["override_value"] is None and cell["effective_value"] == "Delaware"
    assert cell["review_status"] == "approved" and cell["note"] is None
    assert [h["kind"] for h in cell["history"]][-1] == "clear"

    event = next(e for e in audit_events if e["details"].get("action") == "review_cell_reviewed")
    assert event["details"]["cells"] == 1
    assert "New York" not in json.dumps(event["details"])  # counts only


def test_review_validation_and_access(client):
    owner, member, stranger = make_user("owner"), make_user("member"), make_user("stranger")
    matter_id = make_shared_matter(owner, member)
    doc = register_document(owner, matter_id=matter_id)
    tid = _create(client, owner, matter_id=matter_id, document_ids=[doc])["id"]

    # An empty change is refused; a bad status is refused.
    assert _review(client, owner, tid, doc, "governing_law", {}).status_code == 422
    assert (
        _review(client, owner, tid, doc, "governing_law", {"review_status": "maybe"}).status_code
        == 422
    )
    # A cell outside the table's scope does not exist.
    assert _review(client, owner, tid, "nope", "governing_law", {"note": "x"}).status_code == 404
    assert _review(client, owner, tid, doc, "nope", {"note": "x"}).status_code == 404
    # A stranger cannot see the table at all; a member reviews like the owner.
    assert _review(client, stranger, tid, doc, "governing_law", {"note": "x"}).status_code == 404
    resp = _review(client, member, tid, doc, "governing_law", {"review_status": "flagged"})
    assert resp.status_code == 200
    assert resp.json()["review_status"] == "flagged" and resp.json()["reviewed_by"] == member.id
    # Reviewing a cell the model has not reached yet creates it (manual entry).
    assert resp.json()["status"] == "not_found" and resp.json()["value"] is None


# ------------------------------------------------------------------- re-runs


def _changed_answer(value: str) -> dict:
    return {
        "fields": {
            "governing_law": {"value": value, "quote": QUOTE, "not_found": False},
            "effective_date": {"value": "2026-01-05", "quote": "nope", "not_found": False},
        }
    }


def test_rerun_keeps_overrides_and_resets_approval_only_when_the_value_changes(client):
    owner = make_user("owner")
    doc = register_document(owner)
    tid = _create(client, owner, document_ids=[doc])["id"]
    _run_to_completion(client, owner, tid)

    assert (
        _review(client, owner, tid, doc, "governing_law", {"review_status": "approved"}).status_code
        == 200
    )
    assert _review(
        client, owner, tid, doc, "effective_date",
        {"override_value": "2026-02-01", "review_status": "approved"},
    ).status_code == 200  # fmt: skip

    # Same answer again: nothing changed, approvals stand, no "changed" count.
    _started, job, _chat = _run_to_completion(client, owner, tid)
    assert job["result"]["changed"] == 0
    cells = _cells(client, owner, tid)
    assert cells[(doc, "governing_law")]["review_status"] == "approved"
    assert cells[(doc, "effective_date")]["review_status"] == "approved"

    # A different model value: the un-overridden approval drops to pending,
    # the overridden one stands, and both record the model entry in history.
    _started, job, _chat = _run_to_completion(client, owner, tid, answer=_changed_answer("Texas"))
    assert job["result"]["changed"] == 1
    cells = _cells(client, owner, tid)
    gl = cells[(doc, "governing_law")]
    assert gl["value"] == "Texas" and gl["review_status"] == "pending"
    assert [h["kind"] for h in gl["history"]] == ["model", "approve", "model"]
    assert gl["history"][-1]["value"] == "Texas"
    eff = cells[(doc, "effective_date")]
    assert eff["override_value"] == "2026-02-01" and eff["effective_value"] == "2026-02-01"
    assert eff["review_status"] == "approved"  # overridden: the model's change does not reopen it


def test_history_is_capped(client):
    owner = make_user("owner")
    doc = register_document(owner)
    tid = _create(client, owner, document_ids=[doc])["id"]
    _run_to_completion(client, owner, tid)
    for i in range(25):
        assert (
            _review(
                client, owner, tid, doc, "governing_law", {"override_value": f"v{i}"}
            ).status_code
            == 200
        )
    cell = _cells(client, owner, tid)[(doc, "governing_law")]
    assert len(cell["history"]) == 20
    assert cell["history"][-1]["value"] == "v24"


# ---------------------------------------------------------------------- bulk


def test_bulk_review_approves_a_column_and_skips_foreign_refs(client, audit_events):
    owner = make_user("owner")
    d1, d2 = register_document(owner), register_document(owner)
    tid = _create(client, owner, document_ids=[d1, d2])["id"]
    _run_to_completion(client, owner, tid)

    body = {
        "cells": [
            {"document_id": d1, "field_key": "governing_law"},
            {"document_id": d2, "field_key": "governing_law"},
            {"document_id": "nope", "field_key": "governing_law"},
            {"document_id": d1, "field_key": "nope"},
        ],
        "review_status": "approved",
    }
    resp = client.post(f"/api/v1/review-tables/{tid}/cells/bulk", headers=headers(owner), json=body)
    assert resp.status_code == 200, resp.text
    assert resp.json()["updated"] == 2
    assert {c["review_status"] for c in resp.json()["cells"]} == {"approved"}
    cells = _cells(client, owner, tid)
    assert cells[(d1, "governing_law")]["review_status"] == "approved"
    assert cells[(d1, "effective_date")]["review_status"] == "pending"  # other column untouched

    listed = client.get("/api/v1/review-tables", headers=headers(owner)).json()["tables"][0]
    assert (
        listed["approved_count"] == 2 and listed["flagged_count"] == 0 and listed["cell_count"] == 4
    )

    # Nothing to change beyond the refs is refused.
    resp = client.post(
        f"/api/v1/review-tables/{tid}/cells/bulk",
        headers=headers(owner),
        json={"cells": [{"document_id": d1, "field_key": "governing_law"}]},
    )
    assert resp.status_code == 422
    event = next(
        e for e in audit_events if e["details"].get("action") == "review_cells_bulk_reviewed"
    )
    assert event["details"]["cells"] == 2


# ------------------------------------------------------------------- schemas


def test_builtin_and_saved_schemas(client, audit_events):
    owner, member, stranger = make_user("owner"), make_user("member"), make_user("stranger")
    matter_id = make_shared_matter(owner, member)

    listed = client.get("/api/v1/review-tables/schemas", headers=headers(owner)).json()["schemas"]
    builtins = [s for s in listed if s["is_builtin"]]
    assert {s["id"] for s in builtins} == {
        "builtin:nda", "builtin:msa", "builtin:lease", "builtin:employment"
    }  # fmt: skip
    assert all(8 <= len(s["fields"]) <= 12 for s in builtins)
    assert all(f["key"] and f["description"] for s in builtins for f in s["fields"])
    assert all(s["role"] == "builtin" for s in builtins)

    # A personal schema is the owner's alone; a matter schema is shared.
    personal = client.post(
        "/api/v1/review-tables/schemas",
        headers=headers(owner),
        json={"name": "My NDA set", "fields": FIELDS},
    )
    assert personal.status_code == 201, personal.text
    assert personal.json()["role"] == "owner" and personal.json()["is_builtin"] is False
    assert [f["key"] for f in personal.json()["fields"]] == ["governing_law", "effective_date"]
    shared = client.post(
        "/api/v1/review-tables/schemas",
        headers=headers(owner),
        json={"name": "Team set", "matter_id": matter_id, "fields": FIELDS},
    )
    assert shared.status_code == 201
    as_member = client.get("/api/v1/review-tables/schemas", headers=headers(member)).json()[
        "schemas"
    ]
    saved_for_member = {s["name"]: s for s in as_member if not s["is_builtin"]}
    assert set(saved_for_member) == {"Team set"}
    assert saved_for_member["Team set"]["role"] == "member"
    as_stranger = client.get("/api/v1/review-tables/schemas", headers=headers(stranger)).json()
    assert all(s["is_builtin"] for s in as_stranger["schemas"])

    # A matter the caller does not belong to cannot be named.
    resp = client.post(
        "/api/v1/review-tables/schemas",
        headers=headers(stranger),
        json={"name": "x", "matter_id": matter_id, "fields": FIELDS},
    )
    assert resp.status_code == 404

    # Owner edits and deletes; member and built-ins answer 404.
    sid = shared.json()["id"]
    assert (
        client.patch(
            f"/api/v1/review-tables/schemas/{sid}", headers=headers(member), json={"name": "H"}
        ).status_code
        == 404
    )
    assert (
        client.patch(
            "/api/v1/review-tables/schemas/builtin:nda", headers=headers(owner), json={"name": "H"}
        ).status_code
        == 404
    )
    resp = client.patch(
        f"/api/v1/review-tables/schemas/{sid}",
        headers=headers(owner),
        json={"name": "Renamed", "fields": [{"name": "Term", "type": "duration"}]},
    )
    assert resp.status_code == 200 and resp.json()["name"] == "Renamed"
    assert [f["key"] for f in resp.json()["fields"]] == ["term"]
    assert (
        client.patch(
            f"/api/v1/review-tables/schemas/{sid}", headers=headers(owner), json={}
        ).status_code
        == 400
    )
    assert (
        client.delete(f"/api/v1/review-tables/schemas/{sid}", headers=headers(member)).status_code
        == 404
    )
    assert client.delete(f"/api/v1/review-tables/schemas/{sid}", headers=headers(owner)).json() == {
        "status": "deleted",
        "id": sid,
    }
    actions = {e["details"].get("action") for e in audit_events}
    assert {"review_schema_created", "review_schema_updated", "review_schema_deleted"} <= actions


# ----------------------------------------------------------------- documents


def test_add_documents_extracts_only_the_new_rows(client):
    owner, member = make_user("owner"), make_user("member")
    matter_id = make_shared_matter(owner, member)
    d1 = register_document(owner, matter_id=matter_id)
    tid = _create(client, owner, matter_id=matter_id, document_ids=[d1])["id"]
    _run_to_completion(client, owner, tid)

    d2, d3 = register_document(owner, matter_id=matter_id), register_document(owner)
    foreign = register_document(make_user("other"))
    # Member and foreign documents are refused; nothing is added.
    assert (
        client.post(
            f"/api/v1/review-tables/{tid}/documents", headers=headers(member),
            json={"document_ids": [d2]},
        ).status_code
        == 404
    )  # fmt: skip
    assert (
        client.post(
            f"/api/v1/review-tables/{tid}/documents", headers=headers(owner),
            json={"document_ids": [d2, foreign]},
        ).status_code
        == 404
    )  # fmt: skip

    fake_chat = AsyncMock(return_value=_fake_model_response(_answer()))
    with (
        patch.object(ReviewTableService, "_llm_client", return_value=MagicMock()),
        patch.object(review_table_service, "document_text", AsyncMock(return_value=CONTRACT)),
        patch("app.services.review_tables.extraction.openai_chat", fake_chat),
    ):
        resp = client.post(
            f"/api/v1/review-tables/{tid}/documents",
            headers=headers(owner),
            json={"document_ids": [d2, d3, d1], "run": True},
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["added"] == [d2, d3]  # d1 was already in scope
        assert [d["id"] for d in body["table"]["documents"]] == [d1, d2, d3]
        assert body["run"]["documents"] == 2 and body["run"]["run_id"]
        job = _poll(client, owner, body["run"]["job_id"])
    assert job["status"] == "completed"
    assert fake_chat.await_count == 2  # only the new rows
    assert db_count(ReviewCell, ReviewCell.table_id == tid) == 6

    # Removing a document drops its row and cells; an unknown one is 404.
    resp = client.delete(f"/api/v1/review-tables/{tid}/documents/{d2}", headers=headers(owner))
    assert resp.status_code == 200
    assert [d["id"] for d in resp.json()["documents"]] == [d1, d3]
    assert db_count(ReviewCell, ReviewCell.table_id == tid) == 4
    assert (
        client.delete(
            f"/api/v1/review-tables/{tid}/documents/{d2}", headers=headers(owner)
        ).status_code
        == 404
    )
    assert (
        client.delete(
            f"/api/v1/review-tables/{tid}/documents/{d1}", headers=headers(member)
        ).status_code
        == 404
    )


def test_document_text_is_the_text_the_offsets_refer_to(client):
    owner = make_user("owner")
    doc = register_document(owner, filename="nda.pdf")
    other = register_document(owner)
    tid = _create(client, owner, document_ids=[doc])["id"]
    with patch.object(review_table_service, "document_text", AsyncMock(return_value=CONTRACT)):
        resp = client.get(
            f"/api/v1/review-tables/{tid}/documents/{doc}/text", headers=headers(owner)
        )
        assert resp.status_code == 200, resp.text
        assert resp.json() == {
            "document_id": doc, "filename": "nda.pdf", "text": CONTRACT, "length": len(CONTRACT)
        }  # fmt: skip
        assert (
            client.get(
                f"/api/v1/review-tables/{tid}/documents/{other}/text", headers=headers(owner)
            ).status_code
            == 404
        )


# ---------------------------------------------------------------------- runs


def test_run_history_records_scope_counts_and_failures(client):
    owner = make_user("owner")
    d1, d2 = register_document(owner), register_document(owner)
    tid = _create(client, owner, document_ids=[d1, d2])["id"]
    started, _job, _chat = _run_to_completion(client, owner, tid)
    _run_to_completion(
        client, owner, tid, body={"document_ids": [d1]}, answer=_changed_answer("Ohio")
    )

    runs = client.get(f"/api/v1/review-tables/{tid}/runs", headers=headers(owner)).json()["runs"]
    assert len(runs) == 2
    latest, first = runs
    assert first["id"] == started["run_id"] and first["job_id"] == started["job_id"]
    assert first["status"] == "completed" and first["finished_at"]
    assert first["scope"] == {"document_ids": None, "field_keys": None}
    assert first["counts"]["cells"] == 4 and first["counts"]["changed"] == 0
    assert latest["scope"] == {"document_ids": [d1], "field_keys": None}
    assert latest["counts"]["changed"] == 1 and latest["runner_id"] == owner.id

    # A run whose body raises is recorded as failed, with the job.
    with (
        patch.object(ReviewTableService, "_llm_client", return_value=MagicMock()),
        patch.object(review_table_service, "_run", AsyncMock(side_effect=RuntimeError("boom"))),
        patch(
            "app.routers.review_tables.RETRY_POLICIES",
            {"ai_analysis": RetryPolicy(max_retries=0)},
        ),
    ):
        resp = client.post(f"/api/v1/review-tables/{tid}/run", headers=headers(owner), json={})
        assert resp.status_code == 202
        job = _poll(client, owner, resp.json()["job_id"])
    assert job["status"] == "failed"
    failed = client.get(f"/api/v1/review-tables/{tid}/runs", headers=headers(owner)).json()["runs"][
        0
    ]
    assert failed["status"] == "failed" and "boom" in failed["error"]
    assert db_count(ReviewRun, ReviewRun.table_id == tid) == 3

    # Deleting the table takes its runs with it.
    client.delete(f"/api/v1/review-tables/{tid}", headers=headers(owner))
    assert db_count(ReviewRun, ReviewRun.table_id == tid) == 0


# -------------------------------------------------------------------- export


def _export(client, user, tid, **params):
    return client.get(f"/api/v1/review-tables/{tid}/export", headers=headers(user), params=params)


def test_xlsx_export_carries_review_and_summary_sheets(client):
    owner = make_user("owner")
    d1, d2 = register_document(owner, filename="a.pdf"), register_document(owner, filename="b.pdf")
    tid = _create(client, owner, document_ids=[d1, d2])["id"]
    _run_to_completion(client, owner, tid)
    _review(
        client, owner, tid, d1, "governing_law",
        {"override_value": "New York", "review_status": "approved", "note": "ok"},
    )  # fmt: skip
    _review(client, owner, tid, d2, "effective_date", {"review_status": "flagged"})

    resp = _export(client, owner, tid)
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Results", "Citations", "Review", "Summary"]
    results, review, summary = wb["Results"], wb["Review"], wb["Summary"]
    assert results["B2"].value == "New York"  # the effective value
    assert results["B2"].hyperlink is not None
    review_rows = {(r[0], r[1]): r for r in review.iter_rows(min_row=2, values_only=True)}
    assert review_rows[("a.pdf", "Governing law")][2:6] == (
        "Delaware", "New York", "New York", "approved"
    )  # fmt: skip
    assert review_rows[("a.pdf", "Governing law")][6] == owner.id
    assert review_rows[("a.pdf", "Governing law")][8] == "ok"
    assert review_rows[("b.pdf", "Effective date")][5] == "flagged"
    summary_rows = {r[0]: r for r in summary.iter_rows(min_row=2, values_only=True)}
    # Field, verified, unverified, not found, errors, approved, flagged, overridden
    assert summary_rows["Governing law"] == ("Governing law", 2, 0, 0, 0, 1, 0, 1)
    assert summary_rows["Effective date"] == ("Effective date", 0, 2, 0, 0, 0, 1, 0)


def test_export_options_narrow_fields_cells_and_quotes(client, audit_events):
    owner = make_user("owner")
    d1, d2 = register_document(owner, filename="a.pdf"), register_document(owner, filename="b.pdf")
    tid = _create(client, owner, document_ids=[d1, d2])["id"]
    _run_to_completion(client, owner, tid)
    _review(client, owner, tid, d1, "governing_law", {"review_status": "approved"})
    _review(client, owner, tid, d2, "effective_date", {"review_status": "flagged"})

    # fields= keeps only those columns.
    wb = load_workbook(io.BytesIO(_export(client, owner, tid, fields="effective_date").content))
    assert [c.value for c in wb["Results"][1]] == ["Document", "Effective date"]
    assert _export(client, owner, tid, fields="nope").status_code == 400

    # reviewed_only keeps approved cells only.
    wb = load_workbook(io.BytesIO(_export(client, owner, tid, reviewed_only="true").content))
    cites = list(wb["Citations"].iter_rows(min_row=2, values_only=True))
    assert [(r[0], r[1]) for r in cites] == [("a.pdf", "Governing law")]

    # status= filters on extraction or review status.
    wb = load_workbook(io.BytesIO(_export(client, owner, tid, status="flagged").content))
    cites = list(wb["Citations"].iter_rows(min_row=2, values_only=True))
    assert [(r[0], r[1]) for r in cites] == [("b.pdf", "Effective date")]
    assert _export(client, owner, tid, status="bogus").status_code == 400

    # include_quotes=false blanks the quote column everywhere.
    wb = load_workbook(io.BytesIO(_export(client, owner, tid, include_quotes="false").content))
    assert all(r[3] in ("", None) for r in wb["Citations"].iter_rows(min_row=2, values_only=True))
    csv_text = _export(client, owner, tid, format="csv", include_quotes="false").text
    assert QUOTE not in csv_text and "Review status" in csv_text.splitlines()[0]
    assert "approved" in csv_text

    # JSON carries the same cells, their effective values and the summary.
    resp = _export(client, owner, tid, format="json")
    assert resp.status_code == 200 and resp.headers["content-type"].startswith("application/json")
    body = resp.json()
    assert body["table"]["id"] == tid and [f["key"] for f in body["fields"]] == [
        "governing_law", "effective_date"
    ]  # fmt: skip
    assert len(body["cells"]) == 4 and all("effective_value" in c for c in body["cells"])
    assert body["cells"][0]["quote"]
    assert {s["key"]: s["approved"] for s in body["summary"]} == {
        "governing_law": 1, "effective_date": 0
    }  # fmt: skip
    event = [e for e in audit_events if e["details"].get("action") == "review_table_export"][-1]
    assert event["details"]["format"] == "json" and event["details"]["cells"] == 4
